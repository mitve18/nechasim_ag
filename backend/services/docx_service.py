"""
docx_service.py - מנוע הפקת מסמכי Word (RTL) מתבניות מוסדיות.

מה הקובץ עושה:
    ממיר Markdown למסמך Word מימין לשמאל: כותרות, רשימות, **מודגש**, טבלאות
    Markdown (כטבלת Word אמיתית) וציטוטי מקורות (קובץ, סעיף) שהופכים להפניות
    ממוספרות [n] עם סעיף "מקורות" בסוף המסמך.

איך זה עובד:
    - אם קיימת תבנית templates/<doc_type>.docx או templates/base.docx, המסמך נפתח
      ממנה (לוגו, כותרת עליונה/תחתונה וסגנונות מוסדיים נשמרים).
    - מצייני מקום {{field}} בתבנית (כולל בכותרות עליונות/תחתונות) מוחלפים בערכים.
    - גוף המסמך מוזרק במקום הפסקה שמכילה {{body}}; אם אין כזו - בסוף המסמך.
    - בלי תבנית בכלל - מסמך ריק עם גופן Arial ו-RTL (ההתנהגות הקודמת).

תרומה לחיסכון בטוקנים:
    ההמרה מקומית לחלוטין. ייצוא תשובה קיימת ל-Word מהממשק לא עולה אף טוקן,
    והסוכן לא צריך לכתוב מחדש את התוכן כ-argument של כלי.
"""

from __future__ import annotations

import io
import logging
import re
import uuid
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from docx import Document as DocxDocument
from docxtpl import DocxTemplate
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.config import Settings
from backend.services.drive_service import DriveService

logger = logging.getLogger(__name__)

# ציטוט בפורמט שהסוכן מתבקש להשתמש בו: (שם קובץ, סעיף X)
_CITATION_RE = re.compile(r"\(([^()\n]{2,120}?),\s*(סעיף[^()\n]{0,40}|מבוא)\)")
_PLACEHOLDER_RE = re.compile(r"\{\{\s*(\w+)\s*\}\}")
_MD_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")


class DocxService:
    """מפיק מסמכי Word, שומר אותם ב-generated_docs ומגבה לדרייב."""

    def __init__(self, settings: Settings, drive: DriveService,
                 link_resolver: Optional[Callable[[str], Optional[str]]] = None):
        """link_resolver: פונקציה שמחזירה קישור דרייב לפי שם קובץ (לסעיף המקורות)."""
        self.settings = settings
        self.drive = drive
        self.link_resolver = link_resolver
        self.templates_dir = Path(settings.TEMPLATES_DIR)

    # ------------------------------------------------------------------
    # API ציבורי
    # ------------------------------------------------------------------
    def build(self, title: str, body_markdown: str, fields: Optional[Dict[str, str]] = None,
              doc_type: str = "free") -> Dict[str, Any]:
        """
        יוצר מסמך Word ומחזיר {download_url, file_name, drive_link?}.

        fields - ערכים למצייני המקום {{field}} בתבנית. title ו-date נוספים אוטומטית.
        """
        values = {"title": title, "date": datetime.now().strftime("%d/%m/%Y"), **(fields or {})}
        body_markdown, sources = self._extract_citations(body_markdown)

        doc, from_template = self._open_template(doc_type)
        _set_rtl_defaults(doc)
        _replace_placeholders(doc, values, keep={"body"})

        body = doc.element.body
        before = set(id(el) for el in body)
        # בתבנית שכבר מציגה {{title}} לא מוסיפים כותרת כפולה
        if not (from_template and _template_has(doc, "title")):
            _add_paragraph(doc, title, style="Title")
        self._render_markdown(doc, body_markdown)
        if sources:
            self._render_sources(doc, sources)
        _move_new_content_to_anchor(doc, before)

        path = output_path(self.settings, title)
        doc.save(path)
        return publish(self.drive, path)

    def build_from_html(self, title: str, html: str, doc_type: str = "free") -> Dict[str, Any]:
        """
        שלב 2 באשף המסמכים: ממיר את הטיוטה שהמשתמש ערך בממשק (HTML) למסמך Word.

        - התבנית (templates/<doc_type>.docx או base.docx) נטענת עם docxtpl.
        - פסקת {{body}} בתבנית הופכת (בזיכרון) ל-{{p body}}, והגוף מוזרק אליה כ-Subdoc
          שנבנה מה-HTML: כותרות, פסקאות, מודגש/נטוי/קו תחתון, רשימות וטבלאות - הכול RTL.
        - {{title}} ו-{{date}} בתבנית מתמלאים ע"י Jinja של docxtpl.
        - ציטוטי (קובץ, סעיף) הופכים להפניות [n] עם סעיף "מקורות", כמו ב-build.
        ללא LLM - ההמרה מקומית לחלוטין.
        """
        html, sources = self._extract_citations(html)
        tpl = DocxTemplate(self._template_source(doc_type))
        docx = tpl.get_docx()
        _set_rtl_defaults(docx)
        if not _promote_body_anchor(docx):
            docx.add_paragraph("{{p body}}")  # תבנית בלי {{body}} -> הגוף בסוף המסמך

        sd = tpl.new_subdoc()
        if not _template_has(docx, "title"):
            _add_paragraph(sd, title, style="Title")
        parser = _HtmlToDocx(sd)
        parser.feed(html)
        parser.close()
        if sources:
            self._render_sources(sd, sources)

        tpl.render({"title": title, "date": datetime.now().strftime("%d/%m/%Y"), "body": sd})
        path = output_path(self.settings, title)
        tpl.save(path)
        return publish(self.drive, path)

    # ------------------------------------------------------------------
    # תבנית
    # ------------------------------------------------------------------
    def _open_template(self, doc_type: str):
        """פותח templates/<doc_type>.docx, אחרת base.docx, אחרת מסמך ריק."""
        for name in (f"{doc_type}.docx", "base.docx"):
            p = self.templates_dir / name
            if p.is_file():
                try:
                    return DocxDocument(str(p)), True
                except Exception as e:  # noqa: BLE001 - תבנית נעולה/פגומה לא מונעת הפקת מסמך
                    logger.warning("לא ניתן לפתוח את התבנית %s: %s", p.name, e)
        return DocxDocument(), False

    def _template_source(self, doc_type: str) -> Any:
        """
        התבנית ל-docxtpl כ-BytesIO (הקובץ נקרא מיד ולא נשאר נעול).
        תבנית נעולה/פגומה או חסרה -> מסמך ריק בזיכרון, כמו ב-_open_template.
        """
        for name in (f"{doc_type}.docx", "base.docx"):
            p = self.templates_dir / name
            if p.is_file():
                try:
                    data = p.read_bytes()
                    if zipfile.is_zipfile(io.BytesIO(data)):
                        return io.BytesIO(data)
                    logger.warning("התבנית %s אינה קובץ docx תקין", p.name)
                except OSError as e:  # למשל קובץ פתוח ב-Word / נעול ע"י OneDrive
                    logger.warning("לא ניתן לקרוא את התבנית %s: %s", p.name, e)
        buf = io.BytesIO()
        DocxDocument().save(buf)
        buf.seek(0)
        return buf

    # ------------------------------------------------------------------
    # Markdown -> Word
    # ------------------------------------------------------------------
    def _render_markdown(self, doc: Any, md: str) -> None:
        """ממיר Markdown שורה-שורה; רצף שורות שמתחילות ב-| הופך לטבלה."""
        lines = md.splitlines()
        i = 0
        while i < len(lines):
            stripped = lines[i].strip()
            if stripped.startswith("|"):
                block = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    block.append(lines[i].strip())
                    i += 1
                _add_table(doc, block)
                continue
            i += 1
            if not stripped:
                continue
            stripped = _MD_LINK_RE.sub(r"\1", stripped)
            if stripped.startswith("### "):
                _add_paragraph(doc, stripped[4:], style="Heading 3")
            elif stripped.startswith("## "):
                _add_paragraph(doc, stripped[3:], style="Heading 2")
            elif stripped.startswith("# "):
                _add_paragraph(doc, stripped[2:], style="Heading 1")
            elif stripped.startswith(("- ", "* ")):
                _add_paragraph(doc, stripped[2:], style="List Bullet")
            elif re.match(r"^\d+[.)]\s", stripped):
                _add_paragraph(doc, re.sub(r"^\d+[.)]\s", "", stripped), style="List Number")
            else:
                _add_paragraph(doc, stripped)

    # ------------------------------------------------------------------
    # מקורות
    # ------------------------------------------------------------------
    @staticmethod
    def _extract_citations(md: str) -> tuple[str, List[tuple[str, str]]]:
        """מחליף כל (קובץ, סעיף) בהפניה [n] ומחזיר את רשימת המקורות הייחודיים."""
        sources: List[tuple[str, str]] = []

        def repl(m: re.Match) -> str:
            key = (m.group(1).strip(), m.group(2).strip())
            if key not in sources:
                sources.append(key)
            return f"[{sources.index(key) + 1}]"

        return _CITATION_RE.sub(repl, md), sources

    def _render_sources(self, doc: Any, sources: List[tuple[str, str]]) -> None:
        """מוסיף סעיף 'מקורות' ממוספר, עם קישור לקובץ בדרייב כשהוא ידוע."""
        _add_paragraph(doc, "מקורות", style="Heading 2")
        for n, (file_name, section) in enumerate(sources, start=1):
            link = None
            if self.link_resolver:
                try:
                    link = self.link_resolver(file_name)
                except Exception:  # noqa: BLE001
                    link = None
            text = f"[{n}] {file_name}, {section}" + (f" - {link}" if link else "")
            _add_paragraph(doc, text)


# ----------------------------------------------------------------------
# שמירה והחזרת קישור (משותף ל-DocxService ול-TemplateService)
# ----------------------------------------------------------------------
def output_path(settings: Settings, title: str) -> Path:
    """
    נתיב פלט ייחודי ב-generated_docs. השם מכיל רק אותיות/ספרות/עברית/מקף/קו תחתון,
    כך שהוא תמיד עובר את האימות של GET /api/docs/{name} (מניעת Path Traversal).
    """
    safe_title = re.sub(r"[^\w֐-׿\- ]", "", title).strip().replace(" ", "_")[:60] or "document"
    filename = f"{safe_title}_{datetime.now():%Y%m%d_%H%M}_{uuid.uuid4().hex[:6]}.docx"
    return Path(settings.GENERATED_DOCS_DIR) / filename


def publish(drive: DriveService, path: Path) -> Dict[str, Any]:
    """מחזיר קישור הורדה ומגבה לדרייב (Best effort - הדיסק של Render זמני)."""
    result: Dict[str, Any] = {"download_url": f"/api/docs/{path.name}", "file_name": path.name}
    try:
        backup = drive.upload_generated_doc(str(path), path.name)
        if backup:
            result["drive_link"] = backup.get("webViewLink")
    except Exception as e:  # noqa: BLE001 - גיבוי שנכשל לא מבטל את ההורדה המקומית
        logger.warning("גיבוי המסמך לדרייב נכשל: %s", e)
    return result


# ----------------------------------------------------------------------
# פונקציות עזר ל-Word מימין לשמאל
# ----------------------------------------------------------------------
def _set_rtl_defaults(doc: Any) -> None:
    """מגדיר גופן Arial (תומך עברית) כברירת מחדל לכל הסגנונות במסמך."""
    for style in doc.styles:
        try:
            if style.font is not None:
                style.font.name = "Arial"
                rpr = style.element.get_or_add_rPr()
                fonts = rpr.find(qn("w:rFonts"))
                if fonts is None:
                    fonts = OxmlElement("w:rFonts")
                    rpr.append(fonts)
                fonts.set(qn("w:cs"), "Arial")
        except (AttributeError, ValueError):
            continue


def _make_rtl(paragraph: Any) -> None:
    """סימון bidi ברמת הפסקה + יישור לימין."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        ppr.append(bidi)


def _add_runs(paragraph: Any, text: str) -> None:
    """מוסיף טקסט עם תמיכה ב-**מודגש**; כל Run מסומן rtl כדי שעברית תוצג נכון."""
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        run.bold = bold or None
        rpr = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rpr.append(rtl)


def _add_paragraph(doc: Any, text: str, style: Optional[str] = None) -> None:
    """מוסיף פסקה בכיוון RTL (סגנון שלא קיים בתבנית -> פסקה רגילה)."""
    try:
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    except KeyError:
        p = doc.add_paragraph()
    _make_rtl(p)
    _add_runs(p, text)


def _add_table(doc: Any, lines: List[str]) -> None:
    """ממיר טבלת Markdown לטבלת Word עם bidiVisual (עמודה ראשונה מימין)."""
    rows = [
        [c.strip() for c in line.strip().strip("|").split("|")]
        for line in lines if not _TABLE_SEP_RE.match(line.strip())
    ]
    _add_table_rows(doc, rows)


def _add_table_rows(doc: Any, rows: List[List[str]]) -> None:
    """טבלת Word מרשימת שורות (השורה הראשונה מודגשת כשורת כותרת)."""
    rows = [r for r in rows if r]
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = "Table Grid"
    except (KeyError, ValueError):
        pass
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    bidi.set(qn("w:val"), "1")
    tbl_pr.append(bidi)
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            _make_rtl(p)
            text = row[c_idx] if c_idx < len(row) else ""
            _add_runs(p, f"**{text}**" if r_idx == 0 and text else text)
    doc.add_paragraph()  # רווח אחרי הטבלה


def _iter_paragraphs(doc: Any):
    """כל הפסקאות במסמך: גוף, טבלאות, כותרות עליונות ותחתונות."""
    def from_container(container):
        yield from container.paragraphs
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from from_container(cell)

    yield from from_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from from_container(part)


def _template_has(doc: Any, field: str) -> bool:
    return any(f"{{{{{field}}}}}" in p.text.replace(" ", "") for p in _iter_paragraphs(doc))


def _replace_placeholders(doc: Any, values: Dict[str, str], keep: set) -> None:
    """
    מחליף {{field}} בערכים. Word מפצל טקסט ל-Runs באופן שרירותי, ולכן הטקסט
    המלא נכתב ל-Run הראשון ושאר ה-Runs מתרוקנים (העיצוב של ה-Run הראשון נשמר).
    """
    for p in _iter_paragraphs(doc):
        if "{{" not in p.text:
            continue
        new = _PLACEHOLDER_RE.sub(
            lambda m: m.group(0) if m.group(1) in keep else str(values.get(m.group(1), "")), p.text
        )
        if new != p.text and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""


def _promote_body_anchor(doc: Any) -> bool:
    """
    הופך פסקת {{body}} בתבנית ל-{{p body}} - התחביר של docxtpl להזרקת Subdoc
    במקום פסקה שלמה. כך אותה base.docx משמשת גם את build וגם את build_from_html.
    """
    for p in doc.paragraphs:
        if re.search(r"\{\{\s*(?:p\s+)?body\s*\}\}", p.text or ""):
            if p.runs:
                p.runs[0].text = "{{p body}}"
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run("{{p body}}")
            return True
    return False


class _HtmlToDocx(HTMLParser):
    """
    ממיר HTML פשוט (הפלט של עורך הטיוטה בממשק) לפסקאות Word מימין לשמאל.

    נתמך: h1-h4, p/div, br, b/strong, i/em, u, ul/ol/li, table/tr/td/th, blockquote.
    תגיות לא מוכרות - רק הטקסט שבתוכן נשמר; script/style מושמטים לגמרי.
    """

    _BLOCKS = {"p", "div", "h1", "h2", "h3", "h4", "li", "blockquote", "pre"}
    _HEADINGS = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "h4": "Heading 3"}

    def __init__(self, doc: Any):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.runs: List[tuple] = []  # (text, bold, italic, underline)
        self.style: Optional[str] = None
        self.lists: List[str] = []  # מחסנית ul/ol
        self.fmt = {"b": 0, "i": 0, "u": 0}
        self.skip = 0
        self.table: Optional[List[List[str]]] = None
        self.cell: Optional[List[str]] = None

    # --- בלוקים ---
    def _flush(self) -> None:
        runs = [list(r) for r in self.runs]
        # רווחים בקצוות הפסקה (כולל Runs שהם רק רווח/שבירת שורה) מוסרים
        while runs and not runs[0][0].strip():
            runs.pop(0)
        while runs and not runs[-1][0].strip():
            runs.pop()
        if runs:
            runs[0][0] = runs[0][0].lstrip()
            runs[-1][0] = runs[-1][0].rstrip()
            try:
                p = self.doc.add_paragraph(style=self.style) if self.style else self.doc.add_paragraph()
            except KeyError:  # סגנון שלא קיים בתבנית -> פסקה רגילה
                p = self.doc.add_paragraph()
            _make_rtl(p)
            for t, b, it, u in runs:
                if not t:
                    continue
                run = p.add_run(t)
                run.bold = bool(b) or None
                run.italic = bool(it) or None
                run.underline = bool(u) or None
                rtl = OxmlElement("w:rtl")
                rtl.set(qn("w:val"), "1")
                run._r.get_or_add_rPr().append(rtl)
        self.runs = []
        self.style = None

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in ("script", "style"):
            self.skip += 1
        elif tag in ("b", "strong"):
            self.fmt["b"] += 1
        elif tag in ("i", "em"):
            self.fmt["i"] += 1
        elif tag == "u":
            self.fmt["u"] += 1
        elif tag == "br":
            self._text("\n")
        elif tag in ("ul", "ol"):
            self._flush()
            self.lists.append(tag)
        elif tag == "table":
            self._flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.table.append([])
        elif tag in ("td", "th") and self.table is not None:
            self.cell = []
        elif tag in self._BLOCKS:
            self._flush()
            if tag in self._HEADINGS:
                self.style = self._HEADINGS[tag]
            elif tag == "li":
                self.style = "List Number" if self.lists and self.lists[-1] == "ol" else "List Bullet"

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
        elif tag in ("b", "strong"):
            self.fmt["b"] = max(0, self.fmt["b"] - 1)
        elif tag in ("i", "em"):
            self.fmt["i"] = max(0, self.fmt["i"] - 1)
        elif tag == "u":
            self.fmt["u"] = max(0, self.fmt["u"] - 1)
        elif tag in ("td", "th") and self.table is not None and self.cell is not None:
            if not self.table:
                self.table.append([])
            self.table[-1].append(re.sub(r"\s+", " ", "".join(self.cell)).strip())
            self.cell = None
        elif tag == "table" and self.table is not None:
            _add_table_rows(self.doc, self.table)
            self.table = None
        elif tag in ("ul", "ol"):
            self._flush()
            if self.lists:
                self.lists.pop()
        elif tag in self._BLOCKS:
            self._flush()

    def handle_data(self, data: str) -> None:
        if self.skip:
            return
        self._text(re.sub(r"\s+", " ", data))

    def _text(self, text: str) -> None:
        if self.cell is not None:
            self.cell.append(text)
        elif self.table is None:
            self.runs.append((text, self.fmt["b"], self.fmt["i"], self.fmt["u"]))

    def close(self) -> None:
        """סוגר את הפענוח וכותב טקסט אחרון שלא נסגר בתגית (או טבלה שלא נסגרה)."""
        super().close()
        if self.table is not None:
            _add_table_rows(self.doc, self.table)
            self.table = None
        self._flush()


def _move_new_content_to_anchor(doc: Any, before_ids: set) -> None:
    """מעביר את התוכן שנוסף (בסוף הגוף) למיקום הפסקה {{body}} בתבנית, ומוחק אותה."""
    anchor = next((p for p in doc.paragraphs if _PLACEHOLDER_RE.search(p.text or "")
                   and "body" in p.text), None)
    if anchor is None:
        return
    body = doc.element.body
    new_elements = [el for el in body if id(el) not in before_ids and el.tag != qn("w:sectPr")]
    for el in new_elements:
        anchor._p.addprevious(el)
    anchor._p.getparent().remove(anchor._p)

